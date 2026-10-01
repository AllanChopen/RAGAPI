import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document
from fastapi import UploadFile

from app.services.document_service import DocumentService
from app.services.git_service import GitService


class DocxIngestTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.path = Path(self.temp_dir.name) / "especificacion.docx"
        document = Document()
        document.add_paragraph("Requisitos del sistema")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Campo"
        table.cell(0, 1).text = "Tipo"
        table.cell(1, 0).text = "project_id"
        table.cell(1, 1).text = "Texto"
        document.add_paragraph("Consulta de proyectos")
        document.save(self.path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_extracts_paragraphs_and_tables_in_order(self):
        self.assertTrue(GitService.is_supported_filename(self.path.name))
        chunks = GitService.extract_chunks_for_file(
            self.path, self.path.name, chunk_size=1200, chunk_overlap=0, max_chunks=10
        )

        self.assertEqual(len(chunks), 1)
        self.assertEqual(
            chunks[0]["content"].splitlines(),
            [
                "Requisitos del sistema",
                "Campo | Tipo",
                "project_id | Texto",
                "Consulta de proyectos",
            ],
        )
        self.assertEqual(chunks[0]["artifact_type"], "documentation")
        self.assertIsNone(chunks[0]["page"])
        self.assertIsNone(chunks[0]["line_start"])
        self.assertIsNone(chunks[0]["line_end"])

    def test_document_ingest_stores_docx_as_selectable_document(self):
        upload = UploadFile(file=BytesIO(self.path.read_bytes()), filename=self.path.name)
        db = MagicMock()

        with patch.object(
            DocumentService,
            "_embed_in_batches",
            side_effect=lambda texts: [[0.0] * 1536 for _ in texts],
        ):
            result = DocumentService.ingest_files("proyecto-1", [upload], db)

        self.assertEqual(result.total_chunks, 1)
        self.assertEqual(result.documents[0].name, "especificacion.docx")
        self.assertEqual(result.documents[0].artifact_type, "documentation")
        stored = db.add.call_args.args[0]
        self.assertEqual(stored.metadata_json["project_id"], "proyecto-1")
        self.assertEqual(stored.metadata_json["document"], "especificacion.docx")
        self.assertEqual(stored.metadata_json["artifact_type"], "documentation")
        db.commit.assert_called_once()

    def test_invalid_docx_has_no_extractable_content(self):
        self.path.write_bytes(b"not a DOCX archive")
        self.assertEqual(
            GitService.extract_chunks_for_file(
                self.path, self.path.name, chunk_size=1200, chunk_overlap=0, max_chunks=10
            ),
            [],
        )


if __name__ == "__main__":
    unittest.main()
